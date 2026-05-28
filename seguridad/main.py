import webview
import ssl
ssl._create_default_https_context = ssl._create_unverified_context
import os
import sys
import threading
import http.server
import socketserver
import socket
import time
import sqlite3
import openpyxl
import subprocess
from AppKit import NSApp, NSImage

# Importaciones para Selenium
try:
    from selenium import webdriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import Select
    from selenium.webdriver.chrome.options import Options
    import random
except ImportError:
    pass

try:
    from consultar_imei import IMEIScraper as ScraperEstado
    from TomarPantallazo import IMEIScraper as ScraperPantallazo
except ImportError:
    pass

#--íconos en claro u oscuro
def resource_path(relative_path):
    """ Obtiene la ruta absoluta de los recursos (necesario para PyInstaller) """
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)

def set_dock_icon(mode):
    """ Cambia el icono en el Dock según el modo """
    icon_name = "./Vistas/images/logoIMPlight.png" if mode == "light" else "./Vistas/images/logoIMPdark.png"
    icon_path = resource_path(icon_name)
    
    image = NSImage.alloc().initByReferencingFile_(icon_path)
    NSApp.setApplicationIconImage_(image)

def obtener_puerto_libre():
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(('127.0.0.1', 0))
        return s.getsockname()[1]

class Api:
    def __init__(self):
        self.window = None
        self.db_path = self.obtener_ruta_bd()
        self.init_db()

    def obtener_ruta_bd(self):
        home_dir = os.path.expanduser("~")
        app_dir = os.path.join(home_dir, "IMEIManagerData")
        os.makedirs(app_dir, exist_ok=True)
        return os.path.join(app_dir, "imei_database.sqlite")

    def init_db(self):
        self.conn = sqlite3.connect(self.db_path, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        cursor = self.conn.cursor()
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS registros (
                imei TEXT PRIMARY KEY, modelo TEXT, estado TEXT, operador TEXT, 
                cliente TEXT, razon TEXT, encargado TEXT, pago TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS papelera (
                imei TEXT PRIMARY KEY, modelo TEXT, estado TEXT, operador TEXT, 
                cliente TEXT, razon TEXT, encargado TEXT, pago TEXT, 
                fecha_borrado DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        for query in [
            "ALTER TABLE registros ADD COLUMN pin_desbloqueo TEXT DEFAULT ''",
            "ALTER TABLE papelera ADD COLUMN pin_desbloqueo TEXT DEFAULT ''",
            "ALTER TABLE registros ADD COLUMN reg_wom TEXT DEFAULT 'No'",
            "ALTER TABLE papelera ADD COLUMN reg_wom TEXT DEFAULT 'No'",
            "ALTER TABLE registros ADD COLUMN reg_etb TEXT DEFAULT 'No'",
            "ALTER TABLE papelera ADD COLUMN reg_etb TEXT DEFAULT 'No'"
        ]:
            try:
                cursor.execute(query)
            except sqlite3.OperationalError:
                pass

        cursor.execute('''
            CREATE TABLE IF NOT EXISTS encargados (
                nombre TEXT PRIMARY KEY, identificacion TEXT DEFAULT '',
                lugar_expedicion TEXT DEFAULT '', fecha_expedicion TEXT DEFAULT '',
                direccion TEXT DEFAULT '', lineas_wom TEXT DEFAULT '',
                lineas_etb TEXT DEFAULT '', correo TEXT DEFAULT '',
                mensaje TEXT DEFAULT '', color TEXT DEFAULT '#39FF14'
            )
        ''')
        try: cursor.execute("ALTER TABLE encargados ADD COLUMN color TEXT DEFAULT '#39FF14'")
        except sqlite3.OperationalError: pass

        self.conn.commit()

    # --- AUTOMATIZACION ---
    def registrar_wom(self, imei, linea_wom):
        cursor = self.conn.cursor()
        cursor.execute('SELECT encargado FROM registros WHERE imei = ?', (imei,))
        reg = cursor.fetchone()
        
        if not reg or not reg['encargado']: return {"status": "error", "mensaje": "IMEI sin encargado."}
        
        encargado = self.obtener_encargado(reg['encargado'])
        if not encargado: return {"status": "error", "mensaje": "Encargado no encontrado."}

        partes = encargado['nombre'].strip().split()
        primer_nombre = partes[0] if len(partes) > 0 else "NULL"
        segundo_nombre, primer_apellido, segundo_apellido = "NULL", "NULL", "NULL"
        if len(partes) == 2: primer_apellido = partes[1]
        elif len(partes) == 3: segundo_nombre = partes[1]; primer_apellido = partes[2]
        elif len(partes) >= 4: segundo_nombre = partes[1]; primer_apellido = partes[2]; segundo_apellido = " ".join(partes[3:])

        documento = encargado['identificacion']
        if not documento or not linea_wom: return {"status": "error", "mensaje": "Faltan datos de ID o Línea."}

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "RegistrarWom.py")
        
        def _ejecutar_en_background():
            import json
            from datetime import datetime
            try:
                cmd = [sys.executable, script_path, str(imei), str(linea_wom), str(documento), str(primer_nombre), str(segundo_nombre), str(primer_apellido), str(segundo_apellido)]
                proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                stdout, stderr = proc.communicate()
                
                resultado = {"status": "error", "mensaje": "Error de comunicación con Selenium."}
                for line in stdout.splitlines():
                    if line.startswith('{'):
                        try: resultado = json.loads(line)
                        except: pass
                
                if resultado.get("status") == "success":
                    self.actualizar_campo(imei, 'reg_wom', datetime.now().isoformat())
                elif resultado.get("status") == "error":
                    self.actualizar_campo(imei, 'reg_wom', 'Error')
                
                # Enviar alerta a JS y forzar la sincronización correcta
                if self.window:
                    msg = resultado.get('mensaje', '').replace("'", "\\'")
                    self.window.evaluate_js(f"showToast('{msg}', '')")
                    # AQUÍ ESTÁ LA MAGIA: .then() asegura que primero cargue datos y LUEGO actualice los botones
                    self.window.evaluate_js("cargarDatos().then(() => actualizarUIBotonesRegistro())")


            except Exception as e:
                if self.window: self.window.evaluate_js(f"showToast('Error interno: {str(e)}', '')")

        threading.Thread(target=_ejecutar_en_background, daemon=True).start()
        return {"status": "success", "mensaje": "Bot WOM lanzado..."}

    def registrar_etb(self, imei, linea_etb):
        if not linea_etb: return {"status": "error", "mensaje": "Falta la línea ETB."}

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "RegistrarEtb.py")
        
        def _ejecutar_en_background():
            import json
            from datetime import datetime
            try:
                cmd = [sys.executable, script_path, str(imei), str(linea_etb)]
                proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                stdout, stderr = proc.communicate()
                
                resultado = {"status": "error", "mensaje": "Error de comunicación con Selenium ETB."}
                for line in stdout.splitlines():
                    if line.startswith('{'):
                        try: resultado = json.loads(line)
                        except: pass
                
                if resultado.get("status") == "success":
                    self.actualizar_campo(imei, 'reg_etb', datetime.now().isoformat())
                elif resultado.get("status") == "error":
                    self.actualizar_campo(imei, 'reg_etb', 'Error')
                
                # Enviar alerta a JS y forzar la sincronización correcta
                if self.window:
                    msg = resultado.get('mensaje', '').replace("'", "\\'")
                    self.window.evaluate_js(f"showToast('{msg}', '')")
                    # AQUÍ ESTÁ LA MAGIA: .then() asegura que primero cargue datos y LUEGO actualice los botones
                    self.window.evaluate_js("cargarDatos().then(() => actualizarUIBotonesRegistro())")

            except Exception as e:
                if self.window: self.window.evaluate_js(f"showToast('Error interno: {str(e)}', '')")

        threading.Thread(target=_ejecutar_en_background, daemon=True).start()
        return {"status": "success", "mensaje": "Bot ETB lanzado..."}
    def obtener_todos_encargados(self):
        cursor = self.conn.cursor()
        cursor.execute('SELECT * FROM encargados')
        return [dict(fila) for fila in cursor.fetchall()]

    def obtener_encargado(self, nombre):
        cursor = self.conn.cursor()
        cursor.execute('SELECT * FROM encargados WHERE nombre = ?', (nombre,))
        fila = cursor.fetchone()
        if fila:
            return dict(fila)
        return None

    def guardar_encargado(self, datos):
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
                INSERT OR REPLACE INTO encargados 
                (nombre, identificacion, lugar_expedicion, fecha_expedicion, direccion, lineas_wom, lineas_etb, correo, mensaje, color)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                datos['nombre'], datos.get('identificacion', ''), datos.get('lugar_expedicion', ''),
                datos.get('fecha_expedicion', ''), datos.get('direccion', ''), datos.get('lineas_wom', ''),
                datos.get('lineas_etb', ''), datos.get('correo', ''), datos.get('mensaje', ''),
                datos.get('color', '#39FF14')
            ))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def eliminar_encargado(self, nombre):
        try:
            cursor = self.conn.cursor()
            cursor.execute('DELETE FROM encargados WHERE nombre = ?', (nombre,))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def guardar_pin(self, imei, pin):
        try:
            cursor = self.conn.cursor()
            cursor.execute("UPDATE registros SET pin_desbloqueo = ? WHERE imei = ?", (pin, imei))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def set_window(self, window):
        self.window = window

    def obtener_registros(self):
        cursor = self.conn.cursor()
        cursor.execute('SELECT * FROM registros')
        filas = cursor.fetchall()
        return [dict(fila) for fila in filas]

    def guardar_registro(self, datos):
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
                INSERT INTO registros (imei, modelo, estado, operador, cliente, razon, encargado, pago)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ''', (datos['imei'], datos['modelo'], datos['estado'], datos['operador'], 
                  datos['cliente'], datos['razon'], datos['encargado'], datos['pago']))
            self.conn.commit()
            return {"status": "success"}
        except sqlite3.IntegrityError:
            return {"status": "error", "mensaje": "Este IMEI ya está registrado."}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def actualizar_campo(self, imei, campo, valor):
        # ¡AQUÍ ESTABA EL PROBLEMA! Faltaban reg_wom y reg_etb en esta lista:
        campos_permitidos = ["modelo", "cliente", "razon", "encargado", "pago", "reg_wom", "reg_etb"]
        
        if campo not in campos_permitidos:
            return {"status": "error", "mensaje": f"Campo {campo} no modificable de esta forma."}
        
        try:
            cursor = self.conn.cursor()
            query = f"UPDATE registros SET {campo} = ? WHERE imei = ?"
            cursor.execute(query, (valor, imei))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def actualizar_imei(self, imei, headless: bool = False):
        try:
            scraper = ScraperEstado(headless=headless)
            estado, operador = scraper.consultar(imei)
            scraper.close()
            
            cursor = self.conn.cursor()
            cursor.execute('''
                UPDATE registros SET estado = ?, operador = ? WHERE imei = ?
            ''', (estado, operador, imei))
            self.conn.commit()
            
            return {"status": "success", "estado": estado, "operador": operador}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def pantallazo_imei(self, imei, headless: bool = False):
        try:
            home_dir = os.path.expanduser("~")
            app_dir = os.path.join(home_dir, "IMEIManagerData")
            ruta_absoluta = os.path.join(app_dir, f"{imei}.png")
            
            scraper = ScraperPantallazo(headless=headless)
            scraper.consultar_y_capturar(imei, ruta_absoluta)
            scraper.close()
            
            script_osascript = f'set the clipboard to (read (POSIX file "{ruta_absoluta}") as TIFF picture)'
            
            subprocess.run(
                ['/usr/bin/osascript', '-e', script_osascript], 
                stdin=subprocess.DEVNULL, 
                stdout=subprocess.DEVNULL, 
                stderr=subprocess.DEVNULL
            )
            
            if os.path.exists(ruta_absoluta):
                os.remove(ruta_absoluta)
                
            return {"status": "success", "mensaje": "ScreenShot copiado al portapapeles"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def obtener_papelera(self):
        cursor = self.conn.cursor()
        cursor.execute('SELECT * FROM papelera ORDER BY fecha_borrado DESC')
        filas = cursor.fetchall()
        return [dict(fila) for fila in filas]

    def eliminar_registro(self, imei):
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
                INSERT OR REPLACE INTO papelera (imei, modelo, estado, operador, cliente, razon, encargado, pago)
                SELECT imei, modelo, estado, operador, cliente, razon, encargado, pago 
                FROM registros WHERE imei = ?
            ''', (imei,))
            cursor.execute('DELETE FROM registros WHERE imei = ?', (imei,))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def restaurar_registro(self, imei):
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
                INSERT OR REPLACE INTO registros (imei, modelo, estado, operador, cliente, razon, encargado, pago)
                SELECT imei, modelo, estado, operador, cliente, razon, encargado, pago 
                FROM papelera WHERE imei = ?
            ''', (imei,))
            cursor.execute('DELETE FROM papelera WHERE imei = ?', (imei,))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def borrar_permanente(self, imei):
        try:
            cursor = self.conn.cursor()
            cursor.execute('DELETE FROM papelera WHERE imei = ?', (imei,))
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def vaciar_papelera(self):
        try:
            cursor = self.conn.cursor()
            cursor.execute('DELETE FROM papelera')
            self.conn.commit()
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "mensaje": str(e)}

    def consultar_modelos(self, imeis, headless: bool = False):
        import json as _json

        for imei in imeis:
            self.actualizar_campo(imei, 'modelo', 'Consultando...')

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ConsultarModelo.py")

        def _ejecutar_en_background():
            try:
                cmd = [sys.executable, script_path] + [str(i) for i in imeis]
                if headless:
                    cmd.append("--headless")
                proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                for line in proc.stdout:
                    line = line.strip()
                    if not line: continue
                    try:
                        data = _json.loads(line)
                        imei = data.get("imei", "")
                        modelo = data.get("modelo", "")
                        if imei and modelo and not modelo.startswith("Error"):
                            self.actualizar_campo(imei, 'modelo', modelo)
                        elif imei and modelo.startswith("Error"):
                            self.actualizar_campo(imei, 'modelo', 'Error')
                    except _json.JSONDecodeError:
                        pass
                proc.wait()
            except Exception:
                pass

        hilo = threading.Thread(target=_ejecutar_en_background, daemon=True)
        hilo.start()
        return {"status": "success", "mensaje": "Consulta iniciada en segundo plano"}

    def cargar_excel(self):
        try:
            file_types = ('Archivos Excel (*.xlsx;*.xls)', 'Todos los archivos (*.*)')
            # Hack blindado para evitar errores entre versiones antiguas y nuevas de pywebview
            tipo_dialogo = getattr(webview.FileDialog, 'OPEN', getattr(webview, 'OPEN_DIALOG', 1))
            
            result = self.window.create_file_dialog(tipo_dialogo, allow_multiple=False, file_types=file_types)

            if not result:
                return {"status": "cancelled", "mensaje": "Carga cancelada"}
                
            ruta = result[0]
            wb = openpyxl.load_workbook(ruta)
            cursor = self.conn.cursor()
            
            # --- HOJA 1: REGISTROS ---
            sheet_registros = wb.worksheets[0]
            registros_insertados = 0

            for row in sheet_registros.iter_rows(min_row=2, values_only=True):
                row_clean = [str(x) if x is not None else "" for x in row]
                while len(row_clean) < 11: row_clean.append("")

                if row_clean[0].strip() == "": continue

                try:
                    cursor.execute('''
                        INSERT INTO registros (imei, modelo, estado, operador, cliente, razon, encargado, pago, pin_desbloqueo, reg_wom, reg_etb)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (row_clean[0], row_clean[1], row_clean[2], row_clean[3],
                          row_clean[4], row_clean[5], row_clean[6], row_clean[7], 
                          row_clean[8], row_clean[9] or 'No', row_clean[10] or 'No'))
                    registros_insertados += 1
                except sqlite3.IntegrityError:
                    continue

            # --- HOJA 2: ENCARGADOS (OPCIONAL) ---
            encargados_insertados = 0
            if len(wb.sheetnames) > 1:
                sheet_encargados = wb.worksheets[1]
                for row in sheet_encargados.iter_rows(min_row=2, values_only=True):
                    r = [str(x) if x is not None else "" for x in row]
                    while len(r) < 10: r.append("")
                    
                    if r[0].strip() == "": continue
                    try:
                        cursor.execute('''
                            INSERT OR REPLACE INTO encargados 
                            (nombre, identificacion, lugar_expedicion, fecha_expedicion, direccion, lineas_wom, lineas_etb, correo, mensaje, color)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9] or '#39FF14'))
                        encargados_insertados += 1
                    except Exception: pass

            self.conn.commit()
            return {"status": "success", "mensaje": f"Cargados: {registros_insertados} equipos y {encargados_insertados} encargados."}
            
        except Exception as e:
            # Si algo explota en Python, lo atrapamos y lo enviamos a la interfaz de JS
            return {"status": "error", "mensaje": f"Error interno al cargar: {str(e)}"}

    def exportar_excel(self):
        import os
        from datetime import datetime
        try:
            # Apuntar directamente a la carpeta de Descargas del Mac
            home_dir = os.path.expanduser("~")
            descargas_dir = os.path.join(home_dir, "Downloads")
            
            # Crear un nombre único con fecha y hora para no sobreescribir nada
            fecha_str = datetime.now().strftime("%d-%m-%Y_%H-%M-%S")
            nombre_archivo = f"Respaldo_IMEI_{fecha_str}.xlsx"
            ruta_final = os.path.join(descargas_dir, nombre_archivo)
            
            wb = openpyxl.Workbook()
            
            # --- HOJA 1: REGISTROS ---
            ws_reg = wb.active
            ws_reg.title = "Registros"
            headers_reg = ["IMEI", "Modelo", "Estado", "Operador", "Cliente", "Razon", "Encargado", "Pago", "PIN Desbloqueo", "Reg WOM", "Reg ETB"]
            ws_reg.append(headers_reg)
            
            registros = self.obtener_registros()
            for reg in registros:
                ws_reg.append([
                    reg.get('imei', ''), reg.get('modelo', ''), reg.get('estado', ''), reg.get('operador', ''), 
                    reg.get('cliente', ''), reg.get('razon', ''), reg.get('encargado', ''), reg.get('pago', ''), 
                    reg.get('pin_desbloqueo', ''), reg.get('reg_wom', ''), reg.get('reg_etb', '')
                ])
                
            # --- HOJA 2: ENCARGADOS ---
            ws_enc = wb.create_sheet(title="Encargados")
            headers_enc = ["Nombre", "Identificacion", "Lugar Expedicion", "Fecha Expedicion", "Direccion", "Lineas WOM", "Lineas ETB", "Correo", "Mensaje", "Color"]
            ws_enc.append(headers_enc)
            
            encargados = self.obtener_todos_encargados()
            for enc in encargados:
                ws_enc.append([
                    enc.get('nombre', ''), enc.get('identificacion', ''), enc.get('lugar_expedicion', ''), 
                    enc.get('fecha_expedicion', ''), enc.get('direccion', ''), enc.get('lineas_wom', ''), 
                    enc.get('lineas_etb', ''), enc.get('correo', ''), enc.get('mensaje', ''), enc.get('color', '')
                ])
                
            # Guardado directo y forzado
            wb.save(ruta_final)
            
            return {"status": "success", "mensaje": f"¡Guardado en Descargas!"}
            
        except Exception as e:
            return {"status": "error", "mensaje": f"Error al exportar: {str(e)}"}
def iniciar_servidor(ruta_base, puerto):
    class QuietHandler(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, directory=ruta_base, **kwargs)
        def log_message(self, format, *args):
            pass 
            
    socketserver.TCPServer.allow_reuse_address = True
    try:
        with socketserver.TCPServer(("127.0.0.1", puerto), QuietHandler) as httpd:
            httpd.serve_forever()
    except Exception:
        pass

if __name__ == '__main__':
    api = Api()

    if getattr(sys, 'frozen', False):
        application_path = sys._MEIPASS
    else:
        application_path = os.path.dirname(os.path.abspath(__file__))

    ruta_vistas = os.path.join(application_path, "Vistas", "html")
    os.makedirs(ruta_vistas, exist_ok=True)
    ruta_html = os.path.join(ruta_vistas, "Dashboard.html")
    
    if not os.path.exists(ruta_html):
        sys.exit()

    puerto_asignado = obtener_puerto_libre()
    
    hilo_servidor = threading.Thread(target=iniciar_servidor, args=(application_path, puerto_asignado), daemon=True)
    hilo_servidor.start()

    time.sleep(1)

    url_local = f"http://127.0.0.1:{puerto_asignado}/Vistas/html/Dashboard.html"

    ventana = webview.create_window(
        'IMEI Manager Pro',
        url=url_local,
        js_api=api,
        width=1300,
        height=900,
        min_size=(1000, 700),
        background_color='#0f172a'
    )
    
    api.set_window(ventana)
    webview.start()